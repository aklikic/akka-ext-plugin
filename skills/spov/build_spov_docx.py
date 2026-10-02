#!/usr/bin/env python3
"""Build an SPOV scope-of-work DOCX using gdoc-restyle's build_docx, with image support.

    python skills/spov/build_spov_docx.py <blocks.json> <out.docx> --title "<title>" [--meta "Akka · Month Year"]

This wraps build_docx.py from the presentations repo, adding support for
image blocks ({"k": "image", "path": "...", "alt": "..."}) produced by
md_to_blocks.py when rendering Mermaid diagrams.
"""
import argparse
import base64
import json
import os
import sys

# Add gdoc-restyle to path
GDOC_RESTYLE = os.path.expanduser('~/akka/repos/presentations/tools/gdoc-restyle')
sys.path.insert(0, GDOC_RESTYLE)

import build_docx as gd
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu

if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')


def restart_numbering(doc, paragraph):
    """Force an ordered list paragraph to restart numbering at 1.

    The List Number style defines numbering at the style level, so paragraphs
    don't have per-paragraph numPr. We need to:
    1. Find the numId from the style
    2. Create a new numbering instance with a restart override
    3. Add explicit numPr to this paragraph pointing to the new instance
    """
    # Get the numId from the List Number style
    style_el = doc.styles['List Number'].element
    style_pPr = style_el.find(qn('w:pPr'))
    if style_pPr is None:
        return
    style_numPr = style_pPr.find(qn('w:numPr'))
    if style_numPr is None:
        return
    style_numId_el = style_numPr.find(qn('w:numId'))
    if style_numId_el is None:
        return
    style_num_id = style_numId_el.get(qn('w:val'))

    # Access numbering part
    numbering_part = doc.part.numbering_part
    ct_numbering = numbering_part._element

    # Find the abstractNumId for the style's numId
    abstract_id = None
    for num_el in ct_numbering.findall(qn('w:num')):
        if num_el.get(qn('w:numId')) == style_num_id:
            abstract_ref = num_el.find(qn('w:abstractNumId'))
            abstract_id = abstract_ref.get(qn('w:val'))
            break
    if abstract_id is None:
        return

    # Find highest existing numId
    existing = ct_numbering.findall(qn('w:num'))
    max_id = max((int(n.get(qn('w:numId'))) for n in existing), default=0)
    new_id = max_id + 1

    # Create new num with restart override
    new_num = OxmlElement('w:num')
    new_num.set(qn('w:numId'), str(new_id))
    new_abstract = OxmlElement('w:abstractNumId')
    new_abstract.set(qn('w:val'), abstract_id)
    new_num.append(new_abstract)

    lvl_override = OxmlElement('w:lvlOverride')
    lvl_override.set(qn('w:ilvl'), '0')
    start_override = OxmlElement('w:startOverride')
    start_override.set(qn('w:val'), '1')
    lvl_override.append(start_override)
    new_num.append(lvl_override)

    ct_numbering.append(new_num)

    # Add explicit numPr to this paragraph pointing to the new instance
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement('w:numPr')
    ilvl = OxmlElement('w:ilvl')
    ilvl.set(qn('w:val'), '0')
    numId = OxmlElement('w:numId')
    numId.set(qn('w:val'), str(new_id))
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.insert(0, numPr)


def body_with_images(doc, blocks):
    """Extended body() that handles image blocks alongside standard blocks."""
    i = 0
    while i < len(blocks):
        b = blocks[i]
        if b['k'] == 'image':
            # Embed image, scaled to content width
            path = b['path']
            if os.path.exists(path):
                p = doc.add_paragraph()
                gd.para_fmt(p.paragraph_format, before=12, after=12, line=1.0)
                run = p.add_run()
                run.add_picture(path, width=gd.CONTENT_W)
                # Add alt text as caption below
                alt = b.get('alt', '')
                if alt:
                    cap = doc.add_paragraph()
                    gd.para_fmt(cap.paragraph_format, before=0, after=8, line=1.2)
                    gd.set_font(cap.add_run(alt), gd.SANS, 9, gd.MUTED, italic=True)
            else:
                # Image not found — insert placeholder
                p = doc.add_paragraph()
                gd.add_runs(p, [{'t': f'[Image: {path}]', 'b': False, 'i': True}],
                            size=9, color=gd.MUTED)
            i += 1
        elif b['k'] == 'h':
            doc.add_paragraph(b['text'].strip(), style=f"Heading {b['lvl']}")
            i += 1
        elif b['k'] == 'p':
            t = gd.text_of(b['runs']).strip()
            if t.startswith('[Insert > Table of contents'):
                i += 1
            elif b.get('hl') and t.endswith(':'):
                items, i = gd.hyphen_items(blocks, i)
                gd.callout(doc, b['runs'], items, gd.GOLD_DK)
                i += 1
            elif t.startswith('Note:'):
                gd.callout(doc, None, [('para', b['runs'])], gd.ACCENT)
                i += 1
            elif t.startswith('- '):
                items, i = gd.hyphen_items(blocks, i - 1)
                for _, runs in items:
                    gd.add_runs(doc.add_paragraph(style='List Bullet'), runs)
                i += 1
            elif all(x['i'] for x in b['runs']):
                p = doc.add_paragraph()
                gd.para_fmt(p.paragraph_format, before=0, after=6, line=1.4)
                gd.add_runs(p, b['runs'], size=9, color=gd.MUTED)
                i += 1
            else:
                gd.add_runs(doc.add_paragraph(), b['runs'])
                i += 1
        elif b['k'] == 'list':
            style = 'List Number' if b['ordered'] else 'List Bullet'
            for j, it in enumerate(b['items']):
                p = doc.add_paragraph(style=style)
                gd.add_runs(p, it['runs'])
                # Restart numbering at 1 for each ordered list block
                if b['ordered'] and j == 0:
                    restart_numbering(doc, p)
            doc.paragraphs[-1].paragraph_format.space_after = gd.Pt(8)
            i += 1
        elif b['k'] == 'table':
            gd.data_table(doc, b['rows'])
            i += 1
        else:
            i += 1


def build(blocks, title, meta):
    """Build DOCX with image support."""
    from docx import Document
    from docx.shared import Mm

    doc = Document()
    doc.core_properties.title = title
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = gd.MARGIN_LR
    sec.top_margin, sec.bottom_margin = Mm(16), Mm(18)
    sec.footer_distance = Mm(8)
    sec.different_first_page_header_footer = True
    gd.define_styles(doc)
    gd.cover(doc, title, meta)
    body_with_images(doc, blocks)
    gd.footer(sec, title)

    import io
    buf = io.BytesIO()
    doc.save(buf)
    return gd.slim(buf.getvalue())


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('blocks', help='blocks.json from md_to_blocks.py')
    ap.add_argument('out', help='Output .docx path')
    ap.add_argument('--title', required=True, help='Document title (e.g., "Customer — Description")')
    ap.add_argument('--meta', help='Cover meta line (e.g., "Akka · September 2026")')
    args = ap.parse_args()

    blocks = json.load(open(args.blocks, encoding='utf-8'))
    data = build(blocks, args.title, args.meta)

    open(args.out, 'wb').write(data)
    b64 = base64.b64encode(data).decode('ascii')
    open(args.out + '.b64', 'w').write(b64)
    print(f'{args.out}: {len(data)} bytes, {len(b64)} base64 chars')
    if len(b64) > gd.BASE64_WARN:
        print(f'warning: base64 over {gd.BASE64_WARN} chars; the connector upload may be rejected')
